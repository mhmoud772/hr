import csv
import io
from datetime import datetime

from django.core.cache import cache
from apps.attendance.models import Attendance
from shared.logging import get_service_logger
from shared.system import get_settings

logger = get_service_logger("attendance")


class AttendanceService:
    @staticmethod
    def get_attendance_summary(date_param=None):
        qs = Attendance.objects.filter(date=date_param) if date_param else Attendance.objects.all()
        return {
            "total": qs.count(),
            "present": qs.filter(status="present").count(),
            "absent": qs.filter(status="absent").count(),
            "late": qs.filter(status="late").count(),
        }

    @staticmethod
    def close_attendance_day(date_param):
        settings_obj = cache.get("hr_settings_obj") or get_settings(cache_key="hr_settings_obj")
        end_time = None
        if settings_obj:
            end_str = (settings_obj.attendance_settings or {}).get("workEndTime")
            if end_str:
                end_time = datetime.strptime(end_str, "%H:%M").time()

        if not end_time:
            raise ValueError("Work end time not configured.")

        updated = Attendance.objects.filter(
            date=date_param,
            check_out__isnull=True,
            status__in=["present", "late"],
        ).update(check_out=end_time)

        logger.info("Closed attendance day", date=date_param, updated_count=updated)
        return updated

    @staticmethod
    def generate_attendance_report_csv(queryset):
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(
            [
                "Employee Code",
                "Employee Name",
                "Department",
                "Date",
                "Status",
                "Check In",
                "Check Out",
                "Late Minutes",
            ]
        )
        for rec in queryset:
            writer.writerow(
                [
                    rec.employee.employee_code,
                    rec.employee.name,
                    rec.employee.department.name if rec.employee.department else "",
                    rec.date.isoformat(),
                    rec.status,
                    str(rec.check_in or ""),
                    str(rec.check_out or ""),
                    rec.late_minutes or 0,
                ]
            )
        return output.getvalue()

    @staticmethod
    def generate_attendance_report_excel(queryset):
        import openpyxl
        from io import BytesIO
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Report"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="1E3A5F",
            end_color="1E3A5F",
            fill_type="solid",
        )
        headers = [
            "Employee Code",
            "Employee Name",
            "Department",
            "Date",
            "Status",
            "Check In",
            "Check Out",
            "Late Minutes",
        ]

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_idx, rec in enumerate(queryset, 2):
            ws.cell(row=row_idx, column=1, value=rec.employee.employee_code)
            ws.cell(row=row_idx, column=2, value=rec.employee.name)
            ws.cell(
                row=row_idx,
                column=3,
                value=rec.employee.department.name if rec.employee.department else "",
            )
            ws.cell(row=row_idx, column=4, value=rec.date.isoformat())
            ws.cell(row=row_idx, column=5, value=rec.status)
            ws.cell(row=row_idx, column=6, value=str(rec.check_in or ""))
            ws.cell(row=row_idx, column=7, value=str(rec.check_out or ""))
            ws.cell(row=row_idx, column=8, value=rec.late_minutes or 0)

        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = max(12, max_len + 2)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()
